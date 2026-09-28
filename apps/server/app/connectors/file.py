import csv
import io
import json
import zipfile

from openpyxl import load_workbook

from app.connectors.base import BaseConnector

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UPLOAD_ROWS = 5000


def validate_zip_size(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("Spreadsheet archive is too large")
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid spreadsheet archive") from exc


def parse_upload(filename: str, data: bytes) -> list[dict]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File exceeds 5 MB")
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension == "json":
        parsed = json.loads(data.decode("utf-8-sig"))
        rows = parsed if isinstance(parsed, list) else parsed.get("items") if isinstance(parsed, dict) else None
    elif extension == "csv":
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    elif extension == "xlsx":
        validate_zip_size(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("Spreadsheet has no sheet")
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)
            if first is None:
                raise ValueError("Spreadsheet is empty")
            headers = [str(item or "") for item in first]
            rows = []
            for values in iterator:
                rows.append(dict(zip(headers, values)))
                if len(rows) > MAX_UPLOAD_ROWS:
                    raise ValueError("File exceeds 5000 rows")
        finally:
            workbook.close()
    else:
        raise ValueError("Only JSON, CSV and XLSX files are supported")
    if not isinstance(rows, list) or not rows or len(rows) > MAX_UPLOAD_ROWS or any(not isinstance(row, dict) for row in rows):
        raise ValueError("File must contain 1 to 5000 objects")
    return rows


class FileConnector(BaseConnector):
    def __init__(self, filename: str, data: bytes) -> None:
        self.rows = parse_upload(filename, data)

    async def test_connection(self) -> bool:
        return bool(self.rows)

    async def fetch_sample(self) -> list[dict]:
        return self.rows[:5]

    async def fetch_data(self) -> list[dict]:
        return self.rows
